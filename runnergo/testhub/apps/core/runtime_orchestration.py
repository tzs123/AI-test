from __future__ import annotations

import csv
import io
import json
import random
import re
import string
import time
import uuid
from copy import deepcopy
from datetime import datetime, timedelta
from typing import Any, Dict, Iterable, List, Optional, Tuple

from django.utils import timezone
from openpyxl import load_workbook

from .runtime_case import RuntimeCase, normalize_overrides, public_input_fields


TARGET_TYPE_ALIASES = {
    'ui': 'ui_automation',
    'ui_automation': 'ui_automation',
    'app': 'app_automation',
    'app_automation': 'app_automation',
    'api': 'api_automation',
    'api_automation': 'api_automation',
}

_EXPRESSION_RE = re.compile(r'\$\{\s*([^{}]+?)\s*\}')
_TIME_OFFSET_RE = re.compile(r'^time\.now\s*([+-])\s*(\d+)\s*([dhms])$')


def normalize_target_type(value: Any) -> str:
    key = str(value or '').strip().lower()
    target_type = TARGET_TYPE_ALIASES.get(key)
    if not target_type:
        raise ValueError('target_type 必须是 ui_automation、app_automation 或 api_automation')
    return target_type


def _random_phone() -> str:
    return '1' + random.choice('3456789') + ''.join(random.choice(string.digits) for _ in range(9))


def _random_name() -> str:
    last_names = '赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜'
    first_names = '伟刚勇毅俊峰强军平保东文辉力明永健世广志义兴良海山仁波宁贵福生龙元全国胜学祥才发武新利清飞彬富顺信子杰涛昌成康星光天达安岩中茂进林有坚和彪博诚先敬震振壮会思群豪心邦承乐绍功松善厚庆磊民友裕河哲江超浩亮政谦亨奇固之轮翰朗伯宏言若鸣朋斌梁栋维启克伦翔旭鹏泽晨辰士以建家致树炎德行时泰盛雄琛钧冠策腾楠榕风航弘'
    return random.choice(last_names) + ''.join(random.choice(first_names) for _ in range(random.choice([1, 2])))


def _random_string(length: int = 8) -> str:
    length = max(1, min(256, int(length or 8)))
    alphabet = string.ascii_letters + string.digits
    return ''.join(random.choice(alphabet) for _ in range(length))


def _resolve_time_expression(expression: str) -> Optional[str]:
    now = timezone.now()
    if expression == 'time.now':
        return now.isoformat()
    match = _TIME_OFFSET_RE.match(expression)
    if not match:
        return None
    sign, amount_text, unit = match.groups()
    amount = int(amount_text)
    delta_args = {
        'd': {'days': amount},
        'h': {'hours': amount},
        'm': {'minutes': amount},
        's': {'seconds': amount},
    }[unit]
    delta = timedelta(**delta_args)
    value = now + delta if sign == '+' else now - delta
    return value.isoformat()


class RuntimeContext:
    """Runtime-only context for variables, row data and step outputs."""

    def __init__(self, initial: Optional[Dict[str, Any]] = None):
        initial = initial or {}
        self.scopes: Dict[str, Dict[str, Any]] = {
            'global': dict(initial.get('global') or {}),
            'local': dict(initial.get('local') or initial.get('variables') or {}),
            'outputs': dict(initial.get('outputs') or {}),
            'data': dict(initial.get('data') or {}),
            'dataAssets': dict(initial.get('dataAssets') or {}),
            'smartData': dict(initial.get('smartData') or {}),
        }

    def set(self, name: str, value: Any, scope: str = 'local') -> None:
        scope = scope if scope in self.scopes else 'local'
        self.scopes[scope][name] = value

    def update_scope(self, values: Optional[Dict[str, Any]], scope: str = 'data') -> None:
        if not isinstance(values, dict):
            return
        scope = scope if scope in self.scopes else 'data'
        self.scopes[scope].update(values)

    def get(self, name: str, default: Any = None) -> Any:
        name = str(name or '').strip()
        if not name:
            return default
        parts = name.split('.')
        if parts[0] in self.scopes and len(parts) > 1:
            return _get_path(self.scopes.get(parts[0], {}), parts[1:], default)
        for scope in ('local', 'data', 'dataAssets', 'smartData', 'global', 'outputs'):
            if name in self.scopes.get(scope, {}):
                return self.scopes[scope][name]
            value = _get_path(self.scopes.get(scope, {}), parts, None)
            if value is not None:
                return value
        return default

    def as_dict(self) -> Dict[str, Any]:
        return deepcopy(self.scopes)

    def resolve(self, value: Any) -> Any:
        if isinstance(value, str):
            return _EXPRESSION_RE.sub(lambda match: str(self._eval(match.group(1))), value)
        if isinstance(value, list):
            return [self.resolve(item) for item in value]
        if isinstance(value, dict):
            return {key: self.resolve(item) for key, item in value.items()}
        return value

    def _eval(self, expression: str) -> Any:
        expression = str(expression or '').strip()
        lower = expression.lower()
        if lower in {'uuid', 'random.uuid'}:
            return str(uuid.uuid4())
        if lower in {'timestamp', 'time.timestamp'}:
            return str(int(time.time() * 1000))
        if lower in {'phone.random', 'random_phone', 'random_phone_cn'}:
            return _random_phone()
        if lower == 'name.random':
            return _random_name()
        if lower.startswith('random_string'):
            match = re.match(r'random_string(?:\((\d+)\))?$', lower)
            return _random_string(int(match.group(1)) if match and match.group(1) else 8)
        time_value = _resolve_time_expression(lower)
        if time_value is not None:
            return time_value
        value = self.get(expression)
        return value if value is not None else '${' + expression + '}'


def _get_path(source: Any, parts: List[str], default: Any = None) -> Any:
    current = source
    for part in parts:
        if isinstance(current, dict):
            current = current.get(part, default)
        elif isinstance(current, list):
            try:
                current = current[int(part)]
            except (TypeError, ValueError, IndexError):
                return default
        else:
            return default
    return current


def resolve_runtime_value(value: Any, context: Optional[RuntimeContext] = None) -> Any:
    return (context or RuntimeContext()).resolve(value)


def get_case_payload(target_type: str, case_id: int) -> Tuple[Any, str]:
    target_type = normalize_target_type(target_type)
    if target_type in {'app_automation', 'ui_automation'}:
        from apps.app_automation.models import AppTestCase
        test_case = AppTestCase.objects.get(id=case_id)
        return test_case.ui_flow, test_case.name
    from apps.testcases.models import TestCase
    test_case = TestCase.objects.get(id=case_id)
    payload = _parse_case_steps(test_case.steps)
    return payload, test_case.title


def _parse_case_steps(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return value
    if not isinstance(value, str):
        return []
    text = value.strip()
    if not text:
        return []
    try:
        return json.loads(text)
    except (TypeError, ValueError, json.JSONDecodeError):
        return [
            {'id': f'step-{idx}', 'type': 'manual', 'name': line.strip(), 'value': ''}
            for idx, line in enumerate(text.splitlines(), 1)
            if line.strip()
        ]


def scan_case_input_fields(target_type: str, case_id: int) -> Dict[str, Any]:
    payload, case_name = get_case_payload(target_type, case_id)
    return {
        'targetType': normalize_target_type(target_type),
        'caseId': int(case_id),
        'caseName': case_name,
        'fields': public_input_fields(payload),
    }


def parse_dataset_file(file_obj) -> Tuple[List[str], List[Dict[str, Any]]]:
    filename = str(getattr(file_obj, 'name', '') or '').lower()
    raw = file_obj.read()
    if filename.endswith(('.xlsx', '.xlsm')):
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        headers = _normalize_headers(next(rows_iter, []))
        rows = [_row_from_values(headers, values) for values in rows_iter]
    else:
        text = raw.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        raw_headers = list(reader.fieldnames or [])
        headers = _normalize_headers(raw_headers)
        rows = []
        for row in reader:
            normalized_row = {}
            for raw_header, header in zip(raw_headers, headers):
                normalized_row[header] = _normalize_cell(row.get(raw_header))
            rows.append(normalized_row)
    rows = [row for row in rows if any(value not in (None, '') for value in row.values())]
    if not headers:
        raise ValueError('导入文件缺少表头')
    if len(rows) > 1000:
        raise ValueError('单次导入最多支持 1000 组数据')
    return headers, rows


def _normalize_headers(headers: Iterable[Any]) -> List[str]:
    return [str(item or '').strip() for item in headers if str(item or '').strip()]


def _normalize_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if value is None:
        return ''
    return value


def _row_from_values(headers: List[str], values: Iterable[Any]) -> Dict[str, Any]:
    values = list(values or [])
    return {
        header: _normalize_cell(values[index] if index < len(values) else '')
        for index, header in enumerate(headers)
    }


def build_runtime_override_from_row(fields: List[Dict[str, Any]], row: Dict[str, Any]) -> List[Dict[str, Any]]:
    row = row or {}
    result = []
    for field in fields or []:
        value, found = _match_row_value(field, row)
        if not found:
            continue
        result.append({
            'stepId': field.get('stepId'),
            'stepPath': field.get('stepPath'),
            'runtimeValue': value,
        })
    return result


def _match_row_value(field: Dict[str, Any], row: Dict[str, Any]) -> Tuple[Any, bool]:
    keys = [
        field.get('stepId'),
        field.get('stepPath'),
        field.get('step'),
        field.get('element'),
        f"stepId:{field.get('stepId')}",
        f"stepPath:{field.get('stepPath')}",
        f"element:{field.get('element')}",
    ]
    for key in keys:
        if key and key in row:
            return row.get(key), True
    return None, False


def merge_runtime_overrides(*groups: Optional[Iterable[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    merged: Dict[str, Dict[str, Any]] = {}
    order: List[str] = []
    for group in groups:
        for item in normalize_overrides(group or []):
            key = item.get('stepPath') or item.get('stepId')
            if key not in merged:
                order.append(key)
            merged[key] = item
    return [merged[key] for key in order]


def build_runtime_case(
    case_data: Any,
    runtime_override: Optional[Iterable[Dict[str, Any]]] = None,
    context: Optional[RuntimeContext] = None,
) -> RuntimeCase:
    runtime_case = RuntimeCase(case_data, runtime_override).build(context=context)
    return runtime_case


def perform_runtime_cleanup(context: Dict[str, Any], cleanup_config: Optional[Dict[str, Any]]) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """Clean runtime-only context data after execution.

    Supported cleanup_config:
    {
      "clearScopes": ["local"],
      "clearKeys": ["token", "data.phone"],
      "apiRequests": [...]  # recorded as skipped by design
    }
    """
    cleanup_config = cleanup_config or {}
    next_context = deepcopy(context or {})
    result = {
        'cleared': [],
        'skipped': [],
    }

    for scope in cleanup_config.get('clearScopes') or []:
        if scope in next_context and isinstance(next_context[scope], dict):
            next_context[scope] = {}
            result['cleared'].append({'scope': scope})

    for key in cleanup_config.get('clearKeys') or []:
        key = str(key or '').strip()
        if not key:
            continue
        parts = key.split('.')
        if parts[0] in next_context and len(parts) > 1:
            removed = _delete_path(next_context.get(parts[0]), parts[1:])
        else:
            removed = False
            for scope in ('local', 'data', 'dataAssets', 'global', 'outputs'):
                if isinstance(next_context.get(scope), dict) and key in next_context[scope]:
                    del next_context[scope][key]
                    removed = True
                    break
        if removed:
            result['cleared'].append({'key': key})

    for request_config in cleanup_config.get('apiRequests') or []:
        result['skipped'].append({
            'type': 'apiRequest',
            'reason': '外部清理动作需要显式执行器适配，当前仅记录配置',
            'config': request_config,
        })

    return next_context, result


def _delete_path(source: Any, parts: List[str]) -> bool:
    current = source
    for part in parts[:-1]:
        if isinstance(current, dict):
            current = current.get(part)
        elif isinstance(current, list):
            try:
                current = current[int(part)]
            except (TypeError, ValueError, IndexError):
                return False
        else:
            return False
    last = parts[-1]
    if isinstance(current, dict) and last in current:
        del current[last]
        return True
    if isinstance(current, list):
        try:
            current[int(last)] = None
            return True
        except (TypeError, ValueError, IndexError):
            return False
    return False
