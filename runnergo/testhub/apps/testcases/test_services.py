"""
测试用例 Excel 导入服务测试
覆盖 services.py 中 TestCaseImportTemplateService 和 TestCaseExcelImportService 的核心逻辑
"""
import io
from unittest.mock import MagicMock, patch

from django.db import IntegrityError, transaction
from django.test import TestCase as DjangoTestCase, override_settings
from openpyxl import load_workbook

from apps.testcases.models import TestCase, TestCaseImportRecord
from apps.testcases.services import (
    ImportSummary,
    TestCaseDeduplicationService,
    TestCaseExcelImportService,
    TestCaseImportTemplateService,
    INSTRUCTION_SHEET_NAME,
    TEMPLATE_SHEET_NAME,
)
from apps.projects.models import Project
from apps.users.models import User
from apps.versions.models import Version


class TestCaseImportTemplateServiceTests(DjangoTestCase):
    """TestCaseImportTemplateService.build_template() 测试"""

    def test_build_template_returns_bytesio(self):
        """正向：build_template 应返回 BytesIO 对象"""
        result = TestCaseImportTemplateService.build_template()
        self.assertIsInstance(result, io.BytesIO)

    def test_build_template_contains_required_sheet(self):
        """正向：模板应包含 '用例导入模板' 工作表"""
        result = TestCaseImportTemplateService.build_template()
        wb = load_workbook(result)
        self.assertIn(TEMPLATE_SHEET_NAME, wb.sheetnames)

    def test_build_template_contains_instruction_sheet(self):
        """正向：模板应包含 '填写说明' 工作表"""
        result = TestCaseImportTemplateService.build_template()
        wb = load_workbook(result)
        self.assertIn(INSTRUCTION_SHEET_NAME, wb.sheetnames)

    def test_build_template_headers_match_expected(self):
        """正向：模板表头应与 HEADERS 定义一致"""
        result = TestCaseImportTemplateService.build_template()
        wb = load_workbook(result)
        ws = wb[TEMPLATE_SHEET_NAME]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, TestCaseImportTemplateService.HEADERS)

    def test_build_template_contains_sample_row(self):
        """正向：模板应包含示例数据行"""
        result = TestCaseImportTemplateService.build_template()
        wb = load_workbook(result)
        ws = wb[TEMPLATE_SHEET_NAME]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertGreaterEqual(len(rows), 1)
        self.assertEqual(rows[0], tuple(TestCaseImportTemplateService.SAMPLE_ROW))

    def test_build_template_required_headers_have_special_fill(self):
        """边界：必填表头（带*号）应有特殊背景色"""
        result = TestCaseImportTemplateService.build_template()
        wb = load_workbook(result)
        ws = wb[TEMPLATE_SHEET_NAME]
        for index, header in enumerate(TestCaseImportTemplateService.HEADERS, start=1):
            cell = ws.cell(row=1, column=index)
            if header.endswith('*'):
                self.assertIsNotNone(cell.fill.fgColor.rgb)

    def test_build_template_instruction_sheet_has_content(self):
        """正向：填写说明工作表应包含说明内容"""
        result = TestCaseImportTemplateService.build_template()
        wb = load_workbook(result)
        ws = wb[INSTRUCTION_SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        self.assertGreater(len(rows), 0)


class NormalizeTextTests(DjangoTestCase):
    """TestCaseExcelImportService._normalize_text 测试"""

    def test_normalize_text_none_returns_empty(self):
        """边界：None 输入应返回空字符串"""
        self.assertEqual(TestCaseExcelImportService._normalize_text(None), '')

    def test_normalize_text_strips_whitespace(self):
        """正向：应去除首尾空白"""
        self.assertEqual(TestCaseExcelImportService._normalize_text('  hello  '), 'hello')

    def test_normalize_text_number_converts_to_string(self):
        """边界：数字输入应转为字符串"""
        self.assertEqual(TestCaseExcelImportService._normalize_text(123), '123')

    def test_normalize_text_empty_string(self):
        """边界：空字符串应保持"""
        self.assertEqual(TestCaseExcelImportService._normalize_text(''), '')

    def test_normalize_text_preserves_internal_newlines(self):
        """正向：应保留内部换行符"""
        self.assertEqual(TestCaseExcelImportService._normalize_text('line1\nline2'), 'line1\nline2')


class NormalizeHeaderTests(DjangoTestCase):
    """TestCaseExcelImportService._normalize_header 测试"""

    def test_normalize_header_removes_bom(self):
        """边界：应去除 BOM 字符 (\\ufeff)"""
        self.assertEqual(
            TestCaseExcelImportService._normalize_header('\ufeff用例标题*'),
            '用例标题*'
        )

    def test_normalize_header_removes_zero_width(self):
        """边界：应去除零宽字符 (\\u200b)"""
        self.assertEqual(
            TestCaseExcelImportService._normalize_header('用\u200b例标题*'),
            '用例标题*'
        )

    def test_normalize_header_removes_whitespace(self):
        """正向：应去除所有空白字符"""
        self.assertEqual(
            TestCaseExcelImportService._normalize_header(' 用 例 标 题 * '),
            '用例标题*'
        )

    def test_normalize_header_empty_returns_empty(self):
        """边界：空输入应返回空字符串"""
        self.assertEqual(TestCaseExcelImportService._normalize_header(''), '')
        self.assertEqual(TestCaseExcelImportService._normalize_header(None), '')

    def test_normalize_header_combined_special_chars(self):
        """边界：BOM + 零宽 + 空白组合"""
        self.assertEqual(
            TestCaseExcelImportService._normalize_header('\ufeff 用\u200b例标题* '),
            '用例标题*'
        )


class SplitCsvTests(DjangoTestCase):
    """TestCaseExcelImportService._split_csv 测试"""

    def test_split_csv_empty_returns_empty_list(self):
        """边界：空输入返回空列表"""
        self.assertEqual(TestCaseExcelImportService._split_csv(''), [])
        self.assertEqual(TestCaseExcelImportService._split_csv(None), [])

    def test_split_csv_single_value(self):
        """正向：单个值"""
        self.assertEqual(TestCaseExcelImportService._split_csv('V1.0'), ['V1.0'])

    def test_split_csv_multiple_values(self):
        """正向：逗号分隔多个值"""
        self.assertEqual(TestCaseExcelImportService._split_csv('V1.0,V1.1,V2.0'), ['V1.0', 'V1.1', 'V2.0'])

    def test_split_csv_strips_whitespace(self):
        """正向：应去除每项的首尾空白"""
        self.assertEqual(TestCaseExcelImportService._split_csv(' V1.0 , V1.1 '), ['V1.0', 'V1.1'])

    def test_split_csv_ignores_empty_items(self):
        """边界：应忽略空项"""
        self.assertEqual(TestCaseExcelImportService._split_csv('V1.0,,V1.1,'), ['V1.0', 'V1.1'])

    def test_split_csv_number_input(self):
        """边界：数字输入应转为字符串再分割"""
        result = TestCaseExcelImportService._split_csv(123)
        self.assertEqual(result, ['123'])


class ValidateRowTests(DjangoTestCase):
    """TestCaseExcelImportService._validate_row 测试"""

    def test_validate_row_all_required_fields_present(self):
        """正向：所有必填字段存在时无错误"""
        row_data = {
            'title': '登录测试',
            'steps': '1. 打开登录页',
            'expected_result': '登录成功',
        }
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertEqual(errors, [])

    def test_validate_row_missing_title(self):
        """反向：缺少标题应报错"""
        row_data = {'steps': '1. 打开登录页', 'expected_result': '登录成功'}
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertIn('标题不能为空', errors)

    def test_validate_row_missing_steps(self):
        """反向：缺少操作步骤应报错"""
        row_data = {'title': '登录测试', 'expected_result': '登录成功'}
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertIn('操作步骤不能为空', errors)

    def test_validate_row_missing_expected_result(self):
        """反向：缺少预期结果应报错"""
        row_data = {'title': '登录测试', 'steps': '1. 打开登录页'}
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertIn('预期结果不能为空', errors)

    def test_validate_row_all_required_missing(self):
        """反向：所有必填字段缺失应报多个错误"""
        row_data = {}
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertEqual(len(errors), 3)

    def test_validate_row_invalid_priority(self):
        """反向：无效优先级应报错"""
        row_data = {
            'title': '测试',
            'steps': '步骤',
            'expected_result': '结果',
            'priority': 'invalid_priority',
        }
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertTrue(any('优先级无效' in e for e in errors))

    def test_validate_row_valid_priority_chinese(self):
        """正向：中文优先级应通过验证"""
        for priority in ['低', '中', '高', '紧急']:
            row_data = {
                'title': '测试', 'steps': '步骤', 'expected_result': '结果',
                'priority': priority,
            }
            errors = TestCaseExcelImportService._validate_row(row_data, {})
            self.assertEqual(errors, [], f"优先级 '{priority}' 应通过验证")

    def test_validate_row_valid_priority_english(self):
        """正向：英文优先级应通过验证"""
        for priority in ['low', 'medium', 'high', 'critical']:
            row_data = {
                'title': '测试', 'steps': '步骤', 'expected_result': '结果',
                'priority': priority,
            }
            errors = TestCaseExcelImportService._validate_row(row_data, {})
            self.assertEqual(errors, [], f"优先级 '{priority}' 应通过验证")

    def test_validate_row_invalid_test_type(self):
        """反向：无效测试类型应报错"""
        row_data = {
            'title': '测试', 'steps': '步骤', 'expected_result': '结果',
            'test_type': 'invalid_type',
        }
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertTrue(any('测试类型无效' in e for e in errors))

    def test_validate_row_valid_test_type_chinese(self):
        """正向：中文测试类型应通过验证"""
        for test_type in ['功能测试', '集成测试', 'API测试', 'UI测试', '性能测试', '安全测试']:
            row_data = {
                'title': '测试', 'steps': '步骤', 'expected_result': '结果',
                'test_type': test_type,
            }
            errors = TestCaseExcelImportService._validate_row(row_data, {})
            self.assertEqual(errors, [], f"测试类型 '{test_type}' 应通过验证")

    def test_validate_row_valid_test_type_english(self):
        """正向：英文测试类型应通过验证"""
        for test_type in ['functional', 'integration', 'api', 'ui', 'performance', 'security']:
            row_data = {
                'title': '测试', 'steps': '步骤', 'expected_result': '结果',
                'test_type': test_type,
            }
            errors = TestCaseExcelImportService._validate_row(row_data, {})
            self.assertEqual(errors, [], f"测试类型 '{test_type}' 应通过验证")

    def test_validate_row_invalid_version(self):
        """反向：不存在的版本应报错"""
        row_data = {
            'title': '测试', 'steps': '步骤', 'expected_result': '结果',
            'versions': 'V99.0',
        }
        errors = TestCaseExcelImportService._validate_row(row_data, {})
        self.assertTrue(any('版本不存在' in e for e in errors))

    def test_validate_row_valid_version(self):
        """正向：存在的版本应通过验证"""
        version_map = {'V1.0': MagicMock()}
        row_data = {
            'title': '测试', 'steps': '步骤', 'expected_result': '结果',
            'versions': 'V1.0',
        }
        errors = TestCaseExcelImportService._validate_row(row_data, version_map)
        self.assertEqual(errors, [])

    def test_validate_row_mixed_valid_invalid_versions(self):
        """边界：混合有效和无效版本应只报无效版本"""
        version_map = {'V1.0': MagicMock()}
        row_data = {
            'title': '测试', 'steps': '步骤', 'expected_result': '结果',
            'versions': 'V1.0,V99.0',
        }
        errors = TestCaseExcelImportService._validate_row(row_data, version_map)
        self.assertTrue(any('V99.0' in e for e in errors))
        self.assertFalse(any('V1.0' in e for e in errors))


class GenerateImportNoTests(DjangoTestCase):
    """TestCaseExcelImportService.generate_import_no 测试"""

    def test_generate_import_no_starts_with_imp(self):
        """正向：导入编号应以 IMP_ 开头"""
        result = TestCaseExcelImportService.generate_import_no()
        self.assertTrue(result.startswith('IMP_'))

    def test_generate_import_no_contains_timestamp(self):
        """正向：导入编号应包含时间戳"""
        result = TestCaseExcelImportService.generate_import_no()
        # IMP_YYYYMMDDHHMMSS_XXXXXX
        parts = result.split('_')
        self.assertEqual(len(parts), 3)
        self.assertEqual(len(parts[1]), 14)  # YYYYMMDDHHMMSS

    def test_generate_import_no_unique(self):
        """正向：连续生成的编号应不同"""
        no1 = TestCaseExcelImportService.generate_import_no()
        no2 = TestCaseExcelImportService.generate_import_no()
        self.assertNotEqual(no1, no2)


class PriorityMapTests(DjangoTestCase):
    """PRIORITY_MAP 映射完整性测试"""

    def test_priority_map_contains_all_chinese_keys(self):
        """正向：应包含所有中文优先级映射"""
        for key in ['低', '中', '高', '紧急']:
            self.assertIn(key, TestCaseExcelImportService.PRIORITY_MAP)

    def test_priority_map_contains_all_english_keys(self):
        """正向：应包含所有英文优先级映射"""
        for key in ['low', 'medium', 'high', 'critical']:
            self.assertIn(key, TestCaseExcelImportService.PRIORITY_MAP)

    def test_priority_map_values_are_valid_choices(self):
        """正向：映射值应为有效的模型选项"""
        valid_values = {'low', 'medium', 'high', 'critical'}
        for value in TestCaseExcelImportService.PRIORITY_MAP.values():
            self.assertIn(value, valid_values)


class TestTypeMapTests(DjangoTestCase):
    """TEST_TYPE_MAP 映射完整性测试"""

    def test_test_type_map_contains_all_chinese_keys(self):
        """正向：应包含所有中文测试类型映射"""
        for key in ['功能测试', '集成测试', 'API测试', 'UI测试', '性能测试', '安全测试']:
            self.assertIn(key, TestCaseExcelImportService.TEST_TYPE_MAP)

    def test_test_type_map_contains_all_english_keys(self):
        """正向：应包含所有英文测试类型映射"""
        for key in ['functional', 'integration', 'api', 'ui', 'performance', 'security']:
            self.assertIn(key, TestCaseExcelImportService.TEST_TYPE_MAP)


class HeaderAliasesTests(DjangoTestCase):
    """HEADER_ALIASES 映射测试"""

    def test_header_aliases_title_alias(self):
        """正向：'标题*' 应映射到 'title'"""
        self.assertEqual(TestCaseExcelImportService.HEADER_ALIASES.get('标题*'), 'title')

    def test_header_aliases_primary_title(self):
        """正向：'用例标题*' 应映射到 'title'"""
        self.assertEqual(TestCaseExcelImportService.HEADER_ALIASES.get('用例标题*'), 'title')

    def test_expected_header_fields_order(self):
        """正向：EXPECTED_HEADER_FIELDS 顺序应正确"""
        self.assertEqual(
            TestCaseExcelImportService.EXPECTED_HEADER_FIELDS,
            ['title', 'preconditions', 'steps', 'expected_result', 'priority', 'test_type', 'versions']
        )


class ImportRecordIntegrationTests(DjangoTestCase):
    """TestCaseExcelImportService.import_record 集成测试"""

    def setUp(self):
        self.user = User.objects.create_user(
            username='testuser', password='testpass123'
        )
        self.project = Project.objects.create(
            name='测试项目', owner=self.user
        )

    def _create_valid_excel(self, rows_data=None):
        """创建有效的 Excel 导入文件"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = TEMPLATE_SHEET_NAME
        ws.append(TestCaseImportTemplateService.HEADERS)
        if rows_data:
            for row in rows_data:
                ws.append(row)
        else:
            ws.append(TestCaseImportTemplateService.SAMPLE_ROW)
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def _create_record(self, excel_file):
        """创建导入记录"""
        record = TestCaseImportRecord(
            import_no=TestCaseExcelImportService.generate_import_no(),
            project=self.project,
            template_version='v1',
            created_by=self.user,
        )
        record.import_file.save('test_import.xlsx', excel_file, save=True)
        return record

    def test_import_record_missing_template_sheet_raises_error(self):
        """反向：缺少模板工作表应抛出 ValueError"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = '错误的工作表名'
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        record = self._create_record(output)
        with self.assertRaises(ValueError) as ctx:
            TestCaseExcelImportService.import_record(record)
        self.assertIn('缺少工作表', str(ctx.exception))

    def test_import_record_mismatched_headers_raises_error(self):
        """反向：表头不匹配应抛出 ValueError"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = TEMPLATE_SHEET_NAME
        ws.append(['错误表头1', '错误表头2', '错误表头3'])
        ws.append(['数据1', '数据2', '数据3'])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        record = self._create_record(output)
        with self.assertRaises(ValueError) as ctx:
            TestCaseExcelImportService.import_record(record)
        self.assertIn('表头不匹配', str(ctx.exception))

    def test_import_record_valid_data_creates_testcase(self):
        """正向：有效数据应成功创建测试用例"""
        # 使用无版本的自定义数据行，避免版本不存在导致失败
        rows_data = [
            ['登录测试', '已打开登录页', '1. 输入用户名\n2. 输入密码\n3. 点击登录', '登录成功', '高', '功能测试', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.success_count, 1)
        self.assertEqual(summary.failed_count, 0)
        self.assertTrue(TestCase.objects.filter(title='登录测试').exists())

    def test_import_record_empty_rows_returns_zero_total(self):
        """边界：空数据行（仅表头无数据）应返回 total_rows=0"""
        # 不传 rows_data 时模板会包含示例行，需要传空列表
        # 但空列表意味着只有表头，没有数据行
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = TEMPLATE_SHEET_NAME
        ws.append(TestCaseImportTemplateService.HEADERS)
        # 不添加任何数据行
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        record = self._create_record(output)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.total_rows, 0)
        self.assertEqual(summary.success_count, 0)

    def test_import_record_invalid_rows_counted_as_failed(self):
        """反向：无效数据行应计入 failed_count"""
        # 全 None 行会被 effective_rows 过滤掉（空行跳过）
        # 需要使用有内容但必填字段为空的行
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = TEMPLATE_SHEET_NAME
        ws.append(TestCaseImportTemplateService.HEADERS)
        # 只有优先级（非必填），但必填字段为空
        ws.append(['', '', '', '', '高', '', ''])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        record = self._create_record(output)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.failed_count, 1)
        self.assertGreater(len(summary.failure_details), 0)

    def test_import_record_mixed_valid_invalid_rows(self):
        """边界：混合有效和无效行应分别计数"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = TEMPLATE_SHEET_NAME
        ws.append(TestCaseImportTemplateService.HEADERS)
        ws.append(['有效标题', '前置条件', '操作步骤', '预期结果', '高', '功能测试', ''])
        # 有内容但必填字段为空的行（空字符串经过 _normalize_text 后仍为空）
        ws.append(['', '', '', '', '低', '', ''])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        record = self._create_record(output)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.success_count, 1)
        self.assertEqual(summary.failed_count, 1)

    def test_import_record_priority_mapping_chinese(self):
        """正向：中文优先级应正确映射"""
        rows_data = [
            ['高优先级测试', '', '步骤', '结果', '高', '', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.success_count, 1)
        testcase = TestCase.objects.get(title='高优先级测试')
        self.assertEqual(testcase.priority, 'high')

    def test_import_record_test_type_mapping_english(self):
        """正向：英文测试类型应正确映射"""
        rows_data = [
            ['API测试', '', '步骤', '结果', '', 'api', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.success_count, 1)
        testcase = TestCase.objects.get(title='API测试')
        self.assertEqual(testcase.test_type, 'api')

    def test_import_record_default_priority_is_medium(self):
        """边界：未指定优先级时默认为 medium"""
        rows_data = [
            ['默认优先级', '', '步骤', '结果', '', '', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        testcase = TestCase.objects.get(title='默认优先级')
        self.assertEqual(testcase.priority, 'medium')

    def test_import_record_default_test_type_is_functional(self):
        """边界：未指定测试类型时默认为 functional"""
        rows_data = [
            ['默认类型', '', '步骤', '结果', '', '', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        testcase = TestCase.objects.get(title='默认类型')
        self.assertEqual(testcase.test_type, 'functional')

    def test_import_record_status_always_draft(self):
        """正向：导入的用例状态应始终为 draft"""
        rows_data = [
            ['状态测试', '', '步骤', '结果', '', '', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        TestCaseExcelImportService.import_record(record)
        testcase = TestCase.objects.get(title='状态测试')
        self.assertEqual(testcase.status, 'draft')

    def test_import_record_with_valid_version(self):
        """正向：有效版本应正确关联"""
        version = Version.objects.create(name='V1.0', created_by=self.user)
        version.projects.add(self.project)
        rows_data = [
            ['版本测试', '', '步骤', '结果', '', '', 'V1.0'],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.success_count, 1)
        testcase = TestCase.objects.get(title='版本测试')
        self.assertIn(version, testcase.versions.all())

    def test_import_record_with_multiple_versions(self):
        """正向：多个版本应正确关联"""
        v1 = Version.objects.create(name='V1.0', created_by=self.user)
        v1.projects.add(self.project)
        v2 = Version.objects.create(name='V2.0', created_by=self.user)
        v2.projects.add(self.project)
        rows_data = [
            ['多版本测试', '', '步骤', '结果', '', '', 'V1.0,V2.0'],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        testcase = TestCase.objects.get(title='多版本测试')
        self.assertEqual(testcase.versions.count(), 2)

    def test_import_record_with_invalid_version_fails(self):
        """反向：无效版本名应导致行失败"""
        rows_data = [
            ['无效版本测试', '', '步骤', '结果', '', '', 'V99.0'],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        summary = TestCaseExcelImportService.import_record(record)
        self.assertEqual(summary.failed_count, 1)
        self.assertTrue(any('版本不存在' in e.get('errors', [''])[0]
                          for e in summary.failure_details))

    def test_import_record_updates_progress(self):
        """正向：导入过程中应更新进度"""
        rows_data = [
            ['进度测试1', '', '步骤', '结果', '', '', ''],
            ['进度测试2', '', '步骤', '结果', '', '', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        TestCaseExcelImportService.import_record(record)
        record.refresh_from_db()
        self.assertEqual(record.progress, 100)

    def test_import_record_author_is_record_creator(self):
        """正向：导入用例的作者应为导入记录的创建者"""
        rows_data = [
            ['作者测试', '', '步骤', '结果', '', '', ''],
        ]
        excel_file = self._create_valid_excel(rows_data=rows_data)
        record = self._create_record(excel_file)
        TestCaseExcelImportService.import_record(record)
        testcase = TestCase.objects.get(title='作者测试')
        self.assertEqual(testcase.author, self.user)

    def test_import_record_skips_duplicate_titles_in_same_file(self):
        rows_data = [
            ['重复标题', '', '步骤1', '结果1', '', '', ''],
            ['  重复标题  ', '', '步骤2', '结果2', '', '', ''],
        ]
        record = self._create_record(self._create_valid_excel(rows_data=rows_data))

        summary = TestCaseExcelImportService.import_record(record)

        self.assertEqual(summary.success_count, 1)
        self.assertEqual(summary.skip_count, 1)
        self.assertEqual(summary.failed_count, 0)
        self.assertEqual(
            TestCase.objects.filter(project=self.project, title='重复标题').count(),
            1,
        )

    def test_import_record_skips_existing_title_case_insensitively(self):
        TestCase.objects.create(
            project=self.project,
            title='Existing Case',
            steps='原步骤',
            expected_result='原结果',
            author=self.user,
        )
        rows_data = [
            [' existing case ', '', '新步骤', '新结果', '', '', ''],
        ]
        record = self._create_record(self._create_valid_excel(rows_data=rows_data))

        summary = TestCaseExcelImportService.import_record(record)

        self.assertEqual(summary.success_count, 0)
        self.assertEqual(summary.skip_count, 1)
        self.assertEqual(TestCase.objects.filter(project=self.project).count(), 1)


class TestCaseDeduplicationServiceTests(DjangoTestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='dedup-user', password='testpass123'
        )
        self.project = Project.objects.create(name='去重项目', owner=self.user)

    def test_same_title_can_exist_in_different_projects(self):
        other_project = Project.objects.create(name='其他项目', owner=self.user)
        defaults = {
            'steps': '步骤',
            'expected_result': '结果',
            'author': self.user,
        }

        _, first_created = TestCaseDeduplicationService.create(
            project=self.project, title='同名用例', **defaults
        )
        _, second_created = TestCaseDeduplicationService.create(
            project=other_project, title='同名用例', **defaults
        )

        self.assertTrue(first_created)
        self.assertTrue(second_created)

    def test_normalizes_case_width_and_whitespace(self):
        defaults = {
            'steps': '步骤',
            'expected_result': '结果',
            'author': self.user,
        }
        first, first_created = TestCaseDeduplicationService.create(
            project=self.project, title='  Ａgent   Login  ', **defaults
        )
        second, second_created = TestCaseDeduplicationService.create(
            project=self.project, title='agent login', **defaults
        )

        self.assertTrue(first_created)
        self.assertFalse(second_created)
        self.assertEqual(second.id, first.id)
        self.assertEqual(first.title, 'Agent Login')

    def test_database_constraint_blocks_direct_duplicate_create(self):
        TestCase.objects.create(
            project=self.project,
            title='唯一用例',
            steps='步骤',
            expected_result='结果',
            author=self.user,
        )

        with self.assertRaises(IntegrityError), transaction.atomic():
            TestCase.objects.create(
                project=self.project,
                title='  唯一用例  ',
                steps='另一步骤',
                expected_result='另一结果',
                author=self.user,
            )


class BuildFailureReportTests(DjangoTestCase):
    """TestCaseExcelImportService.build_failure_report 测试"""

    def test_build_failure_report_no_failures_returns_none(self):
        """正向：无失败记录时返回 None"""
        record = MagicMock()
        record.failure_details = []
        result = TestCaseExcelImportService.build_failure_report(record)
        self.assertIsNone(result)

    def test_build_failure_report_with_failures_returns_content_file(self):
        """正向：有失败记录时返回 ContentFile"""
        record = MagicMock()
        record.failure_details = [
            {'row_number': 2, 'title': '测试', 'errors': ['标题不能为空']}
        ]
        record.import_no = 'IMP_20240101000000_ABCDEF'
        result = TestCaseExcelImportService.build_failure_report(record)
        self.assertIsNotNone(result)
        self.assertTrue(result.name.endswith('.xlsx'))

    def test_build_failure_report_contains_error_details(self):
        """正向：失败报告应包含错误详情"""
        record = MagicMock()
        record.failure_details = [
            {'row_number': 2, 'title': '测试用例', 'errors': ['标题不能为空', '步骤不能为空']}
        ]
        record.import_no = 'IMP_20240101000000_ABCDEF'
        result = TestCaseExcelImportService.build_failure_report(record)
        wb = load_workbook(io.BytesIO(result.read()))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 2)  # 行号
        self.assertEqual(rows[0][1], '测试用例')  # 标题


class ImportSummaryDataclassTests(DjangoTestCase):
    """ImportSummary 数据类测试"""

    def test_import_summary_creation(self):
        """正向：应正确创建 ImportSummary"""
        summary = ImportSummary(
            total_rows=10,
            success_count=8,
            failed_count=2,
            skip_count=0,
            failure_details=[{'row_number': 3, 'errors': ['错误']}]
        )
        self.assertEqual(summary.total_rows, 10)
        self.assertEqual(summary.success_count, 8)
        self.assertEqual(summary.failed_count, 2)
        self.assertEqual(summary.skip_count, 0)
        self.assertEqual(len(summary.failure_details), 1)
